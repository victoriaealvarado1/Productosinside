#!/usr/bin/env python3
"""Lee una matriz de Payless (.pptx) y la deja en un JSON comparable.

Uso:  python3 leer_matriz.py MATRIZ.pptx SALIDA_DIR

Escribe SALIDA_DIR/<nombre>.json y guarda las artes y videos de cada diapositiva
en SALIDA_DIR/media/<sha1>.<ext>. Cada diapositiva se identifica por su id interno
de PowerPoint (no cambia si el cliente inserta o mueve diapositivas) y cada
comentario por un id estable calculado con autor, fecha y texto.
"""
import hashlib, json, os, sys, zipfile
from pptx import Presentation
from lxml import etree

NS = {
    "p188": "http://schemas.microsoft.com/office/powerpoint/2018/8/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}
ETIQUETAS = {"FORMATO": "formato", "NOMBRE PIEZA": "nombre", "CATEGORIA CONTENIDO": "categoria", "PLATAFORMA": "plataforma",
             "FECHA PUBLICACIÓN": "fecha", "FECHA DE PUBLICACIÓN": "fecha", "SERIE / TIPO": "serie", "VIGENCIA": "vigencia"}
MIN_KB = 60  # imágenes más chicas son logos e íconos de la plantilla


def sha1(b): return hashlib.sha1(b).hexdigest()


def cid(autor, fecha, texto):
    return sha1(f"{autor}|{fecha}|{texto}".encode("utf-8"))[:12]


def archivo_sha1(path):
    h = hashlib.sha1()
    with open(path, "rb") as f:
        for blk in iter(lambda: f.read(1 << 20), b""): h.update(blk)
    return h.hexdigest()


def leer(path, salida):
    prs = Presentation(path); W, H = prs.slide_width, prs.slide_height
    z = zipfile.ZipFile(path)
    autores = {}
    if "ppt/authors.xml" in z.namelist():
        for a in etree.fromstring(z.read("ppt/authors.xml")).findall("p188:author", NS): autores[a.get("id")] = a.get("name")
    media = os.path.join(salida, "media"); os.makedirs(media, exist_ok=True)
    slides = []
    for n, s in enumerate(prs.slides, 1):
        sl = {"id": str(s.slide_id), "n": n, "textos": [], "campos": {}, "artes": [], "videos": [], "comentarios": []}
        orden = [0]

        def guardar(part):
            b = part.blob; h = sha1(b); ext = os.path.splitext(part.partname)[1] or ".bin"
            fn = os.path.join(media, h + ext)
            if not os.path.exists(fn): open(fn, "wb").write(b)
            return h, ext, len(b) // 1024

        def caminar(shapes):
            for sh in shapes:
                if sh.shape_type == 6 and hasattr(sh, "shapes"): caminar(sh.shapes); continue
                L = sh.left or 0; T = sh.top or 0; w = sh.width or 0; h = sh.height or 0
                dentro = L < W and T < H and L + w > 0 and T + h > 0
                pos = {"x": round(L / W, 3), "y": round(T / H, 3), "w": round(w / W, 3), "h": round(h / H, 3)}
                if sh.has_text_frame and sh.text_frame.text.strip() and dentro:
                    sl["textos"].append({"t": sh.text_frame.text.strip(), **pos})
                if getattr(sh, "has_table", False) and sh.has_table and dentro:
                    for row in sh.table.rows:
                        for c in row.cells:
                            if c.text.strip(): sl["textos"].append({"t": c.text.strip(), **pos})
                el = sh._element
                vids = el.findall(".//a:videoFile", NS)
                for v in vids:
                    try:
                        hh, ext, kb = guardar(s.part.related_part(v.get("{%s}link" % NS["r"])))
                        sl["videos"].append({"sha1": hh, "ext": ext, "kb": kb, "dentro": dentro, **pos})
                    except Exception: pass
                for b in el.findall(".//a:blip", NS):
                    rid = b.get("{%s}embed" % NS["r"])
                    if not rid: continue
                    try: part = s.part.related_part(rid)
                    except Exception: continue
                    orden[0] += 1
                    if len(part.blob) // 1024 < MIN_KB: continue
                    hh, ext, kb = guardar(part)
                    sl["artes"].append({"sha1": hh, "ext": ext, "kb": kb, "capa": orden[0], "dentro": dentro, "poster": bool(vids), **pos})

        caminar(s.shapes)
        # campos de la ficha: etiqueta a la izquierda, valor a la derecha en la misma línea
        for t in sl["textos"]:
            k = ETIQUETAS.get(t["t"].strip().upper())
            if not k: continue
            val = [u for u in sl["textos"] if u is not t and abs(u["y"] - t["y"]) < 0.03 and u["x"] > t["x"] + 0.05]
            if val: sl["campos"][k] = min(val, key=lambda u: u["x"])["t"]
        paises = [u for u in sl["textos"] if u["x"] < 0.12 and u["y"] < 0.08]
        if paises: sl["campos"]["paises"] = paises[0]["t"]
        sl["huella"] = sha1("\n".join(sorted(t["t"] for t in sl["textos"])).encode("utf-8"))[:12]
        for rel in s.part.rels.values():
            if "comments" not in rel.reltype: continue
            x = etree.fromstring(rel.target_part.blob)
            txt = lambda e: " ".join(t.text or "" for t in e.findall("p188:txBody//a:t", NS)).strip()
            for cm in x.findall("p188:cm", NS):
                a, f, t = autores.get(cm.get("authorId"), cm.get("authorId")), cm.get("created"), txt(cm)
                top = cid(a, f, t)
                sl["comentarios"].append({"cid": top, "padre": None, "autor": a, "fecha": f, "texto": t, "resuelto": cm.get("status") == "resolved"})
                for r in cm.findall("p188:replyLst/p188:reply", NS):
                    ra, rf, rt = autores.get(r.get("authorId")), r.get("created"), txt(r)
                    sl["comentarios"].append({"cid": cid(ra, rf, rt), "padre": top, "autor": ra, "fecha": rf, "texto": rt, "resuelto": False})
        slides.append(sl)
    return {"archivo": os.path.basename(path), "sha1": archivo_sha1(path), "slides": slides}


if __name__ == "__main__":
    if len(sys.argv) != 3: sys.exit(__doc__)
    path, salida = sys.argv[1], sys.argv[2]
    os.makedirs(salida, exist_ok=True)
    d = leer(path, salida)
    out = os.path.join(salida, os.path.splitext(os.path.basename(path))[0] + ".json")
    json.dump(d, open(out, "w"), ensure_ascii=False, indent=1)
    print(f"{d['archivo']}: {len(d['slides'])} diapositivas, {sum(len(s['comentarios']) for s in d['slides'])} comentarios -> {out}")
